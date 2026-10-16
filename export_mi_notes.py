# -*- coding: utf-8 -*-
"""
小米/红米笔记 批量导出并合并为一个 Word（.docx）
=================================================

原理：
  红米笔记开启「云同步」后，会同步到网页版小米云 i.mi.com。
  登录后凭 Cookie 调用其网页内部接口，即可批量拉取普通笔记全文与图片：
    列表：GET https://i.mi.com/note/full/page/?ts=毫秒&limit=200[&syncTag=...]
    详情：GET https://i.mi.com/note/note/{笔记ID}/?ts=毫秒
    附件：GET https://i.mi.com/file/full?type=note_img&fileid={附件ID}&ts=毫秒
  本脚本只读取数据，不会修改、删除云端任何内容。

用法：
  1) 先在同目录 cookie.txt 里粘贴登录后的 Cookie（见《使用说明.md》）
  2) python export_mi_notes.py
  3) 不登录先看排版效果：python export_mi_notes.py --demo

依赖：pip install -r requirements.txt
"""

import os
import re
import sys
import json
import time
import base64
import argparse
from html.parser import HTMLParser
from datetime import datetime

try:
    import requests
except ImportError:
    requests = None  # --from-json 模式无需联网，不依赖 requests

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("缺少依赖 python-docx，请先运行：pip install -r requirements.txt")
    sys.exit(1)

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


# --------------------------------------------------------------------------- #
# 配置
# --------------------------------------------------------------------------- #
BASE = "https://i.mi.com"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
COOKIE_FILE = os.path.join(SCRIPT_DIR, "cookie.txt")
ASSETS_DIR = os.path.join(SCRIPT_DIR, "assets")

PAGE_LIMIT = 200          # 每页条数（接口上限 200）
REQUEST_DELAY = 0.35      # 每次请求间隔秒数，避免过快
MAX_IMG_WIDTH_CM = 15.0   # 图片在 Word 中的最大宽度
SORT_BY_MODIFY = False    # False=按创建时间升序；True=按修改时间排序

HEADERS = {
    "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"),
    "Referer": "https://i.mi.com/note/h5",
    "Origin": "https://i.mi.com",
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Accept-Encoding": "gzip, deflate",
    "Sec-Fetch-Dest": "empty",
    "Sec-Fetch-Mode": "cors",
    "Sec-Fetch-Site": "same-origin",
}


# --------------------------------------------------------------------------- #
# 工具函数
# --------------------------------------------------------------------------- #
def now_ms():
    return int(time.time() * 1000)


def ts_to_str(ms):
    """毫秒时间戳 -> 本地时间字符串"""
    try:
        return datetime.fromtimestamp(int(ms) / 1000).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return "未知时间"


def safe_filename(name, limit=80):
    name = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", str(name)).strip().strip(".")
    return (name or "未命名")[:limit]


def bgr_to_rgb(color):
    """小米背景色是 BGR 的 6 位十六进制，转 RGB"""
    if not color:
        return None
    h = color.replace("#", "")
    if len(h) < 6:
        return None
    b, g, r = h[0:2], h[2:4], h[4:6]
    return f"#{r}{g}{b}".upper()


MIME_EXT = {
    "image/jpeg": "jpg", "image/jpg": "jpg", "image/png": "png",
    "image/gif": "gif", "image/webp": "webp", "image/bmp": "bmp",
    "audio/aac": "aac", "audio/mp4": "m4a", "audio/mpeg": "mp3",
    "audio/ogg": "ogg", "audio/x-m4a": "m4a",
    "video/mp4": "mp4", "video/3gpp": "3gp",
}


def ext_from_mime(mime, default="jpg"):
    if not mime:
        return default
    return MIME_EXT.get(mime.split(";")[0].strip().lower(),
                        mime.split("/")[-1].lower() if "/" in mime else default)


# --------------------------------------------------------------------------- #
# 小米云接口
# --------------------------------------------------------------------------- #
class MiNoteClient:
    def __init__(self, cookie):
        self.session = requests.Session()
        self.session.headers.update(HEADERS)
        self.session.headers["Cookie"] = cookie
        self._media_cache = {}

    def _get_json(self, url, retry=2):
        for attempt in range(retry + 1):
            try:
                resp = self.session.get(url, timeout=30)
                if resp.status_code == 401:
                    raise RuntimeError("Cookie 已过期（401），请重新复制 Cookie")
                if resp.status_code == 200:
                    return resp.json()
                time.sleep(0.5 * (attempt + 1))
            except RuntimeError:
                raise
            except Exception as e:
                if attempt == retry:
                    raise RuntimeError(f"请求失败：{url} ({e})")
                time.sleep(0.5 * (attempt + 1))
        raise RuntimeError(f"请求异常：{url}")

    def validate(self):
        url = f"{BASE}/note/full/page/?ts={now_ms()}&limit=1"
        data = self._get_json(url)
        return data.get("result") == "ok"

    def fetch_all_entries(self):
        """翻页拉取全部笔记列表 + 文件夹"""
        entries, folders = [], {}
        sync_tag = ""
        page = 0
        while True:
            page += 1
            url = f"{BASE}/note/full/page/?ts={now_ms()}&limit={PAGE_LIMIT}"
            if sync_tag:
                url += f"&syncTag={sync_tag}"
            data = self._get_json(url)
            if data.get("result") != "ok" or not data.get("data"):
                raise RuntimeError(f"列表接口返回异常：{json.dumps(data, ensure_ascii=False)[:300]}")
            block = data["data"]
            for f in block.get("folders", []) or []:
                folders[str(f.get("id"))] = (f.get("subject") or "未命名文件夹").strip()
            batch = block.get("entries", []) or []
            entries.extend(batch)
            print(f"\r  列表第 {page} 页，累计 {len(entries)} 条", end="", flush=True)
            sync_tag = block.get("syncTag") or sync_tag
            if block.get("lastPage") or not block.get("syncTag"):
                break
            time.sleep(REQUEST_DELAY)
        print()
        return entries, folders

    def fetch_detail(self, note_id):
        url = f"{BASE}/note/note/{note_id}/?ts={now_ms()}"
        data = self._get_json(url)
        if data.get("result") == "ok" and data.get("data", {}).get("entry"):
            return data["data"]["entry"]
        return None

    def download_media(self, raw_id, mime):
        """下载附件，返回 (本地路径或None, 类型img/audio/video)"""
        if not raw_id:
            return None, None
        if raw_id in self._media_cache:
            return self._media_cache[raw_id]
        file_id = raw_id.split(".", 1)[1] if "." in raw_id else raw_id
        category = (mime or "").split("/")[0].lower()
        kind = category if category in ("image", "audio", "video") else "image"
        ext = ext_from_mime(mime, "jpg" if kind == "image" else "bin")
        local_name = f"{kind}_{file_id[-12:]}.{ext}"
        local_path = os.path.join(ASSETS_DIR, local_name)
        if not os.path.exists(local_path):
            url = f"{BASE}/file/full?type=note_img&fileid={file_id}&ts={now_ms()}"
            try:
                resp = self.session.get(url, timeout=60, allow_redirects=True)
                if resp.status_code == 200 and resp.content:
                    os.makedirs(ASSETS_DIR, exist_ok=True)
                    with open(local_path, "wb") as fp:
                        fp.write(resp.content)
                else:
                    print(f"\n  [警告] 附件下载失败 {file_id} HTTP {resp.status_code}")
                    local_path = None
            except Exception as e:
                print(f"\n  [警告] 附件下载异常 {file_id}: {e}")
                local_path = None
        result_kind = "img" if kind == "image" else kind
        self._media_cache[raw_id] = (local_path, result_kind)
        return local_path, result_kind


# --------------------------------------------------------------------------- #
# 笔记解析
# --------------------------------------------------------------------------- #
def parse_extra_info(raw):
    if not raw:
        return {}
    if isinstance(raw, dict):
        return raw
    try:
        return json.loads(raw)
    except Exception:
        return {}


def extract_files(entry):
    """返回 {raw_id: mime}"""
    file_map = {}
    setting = entry.get("setting")
    if isinstance(setting, str):
        try:
            setting = json.loads(setting)
        except Exception:
            setting = None
    raw_files = []
    if isinstance(setting, dict) and isinstance(setting.get("data"), list):
        raw_files = setting["data"]
    elif isinstance(entry.get("files"), list):
        raw_files = entry["files"]
    for f in raw_files:
        rid = f.get("rawId") or f.get("fileId") or f.get("digest")
        if rid:
            file_map[rid] = f.get("mimeType", "")
    return file_map


def normalize_note(entry, detail, folders):
    """合并列表与详情，产出标准 note 字典"""
    merged = dict(entry)
    if detail:
        merged.update(detail)
    extra = parse_extra_info(merged.get("extraInfo"))
    # 思维导图笔记跳过（本工具只处理普通图文）
    if extra.get("note_content_type") == "mind" or extra.get("mind_content"):
        return None
    content = merged.get("content") or merged.get("snippet") or ""
    if not content:
        return None
    title = (extra.get("title") or merged.get("subject") or "").strip()
    if not title:
        title = ts_to_str(merged.get("createDate") or now_ms())
    folder_id = str(merged.get("folderId") or "")
    folder_name = folders.get(folder_id) or "未分组"
    return {
        "id": str(merged.get("id")),
        "title": title,
        "folder_id": folder_id,
        "folder_name": folder_name,
        "content": content,
        "files": extract_files(merged),
        "create": merged.get("createDate") or 0,
        "modify": merged.get("modifyDate") or 0,
    }


# --------------------------------------------------------------------------- #
# 行内样式解析（<b><i><u><delete><background color> 等）
# --------------------------------------------------------------------------- #
class InlineRenderer(HTMLParser):
    def __init__(self, paragraph):
        super().__init__(convert_charrefs=True)
        self.p = paragraph
        self.stack = []  # 每元素一个样式 dict

    def _cur_style(self):
        style = {"bold": False, "italic": False, "underline": False,
                 "strike": False, "fill": None}
        for s in self.stack:
            for k, v in s.items():
                if v:
                    style[k] = v
        return style

    def handle_starttag(self, tag, attrs):
        ad = dict(attrs)
        st = {"bold": False, "italic": False, "underline": False,
              "strike": False, "fill": None}
        if tag == "b":
            st["bold"] = True
        elif tag == "i":
            st["italic"] = True
        elif tag == "u":
            st["underline"] = True
        elif tag == "delete":
            st["strike"] = True
        elif tag in ("size", "mid-size", "h3-size"):
            st["bold"] = True  # 行内出现的标题标签按粗体处理
        elif tag == "background":
            st["fill"] = bgr_to_rgb(ad.get("color", ""))
        elif tag == "br":
            self.p.add_run().add_break()
            return
        self.stack.append(st)

    def handle_endtag(self, tag):
        close = {"b": "bold", "i": "italic", "u": "underline",
                 "delete": "strike", "background": "fill",
                 "size": "bold", "mid-size": "bold", "h3-size": "bold"}
        if tag in close:
            for i in range(len(self.stack) - 1, -1, -1):
                if self.stack[i].get(close[tag]):
                    self.stack.pop(i)
                    break

    def handle_data(self, data):
        if not data:
            return
        st = self._cur_style()
        run = self.p.add_run(data)
        run.bold = st["bold"] or None
        run.italic = st["italic"] or None
        run.underline = st["underline"] or None
        run.font.strike = st["strike"] or None
        if st["fill"]:
            self._set_shading(run, st["fill"])

    @staticmethod
    def _set_shading(run, hex_rgb):
        rpr = run._r.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_rgb.lstrip("#"))
        rpr.append(shd)


MEDIA_RE = re.compile(r"<(img|sound|video)\b[^>]*?/?>")
MEDIA_ID_RE = re.compile(r"(?:id|data)\s*=\s*[\"']([^\"']*)[\"']")


def split_media(text):
    """把一行切成 [('text', 片段), ('img'/'sound'/'video', raw_id), ...]"""
    tokens = []
    pos = 0
    for m in MEDIA_RE.finditer(text):
        if m.start() > pos:
            tokens.append(("text", text[pos:m.start()]))
        idm = MEDIA_ID_RE.search(m.group(0))
        tokens.append((m.group(1), idm.group(1) if idm else None))
        pos = m.end()
    if pos < len(text):
        tokens.append(("text", text[pos:]))
    return tokens


# --------------------------------------------------------------------------- #
# 渲染成 Word
# --------------------------------------------------------------------------- #
class DocBuilder:
    def __init__(self, client=None):
        self.client = client
        self.doc = Document()
        self._setup_styles()
        self.order_counter = 0

    def _setup_styles(self):
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        for lvl, size in ((1, 20), (2, 16), (3, 14)):
            st = self.doc.styles[f"Heading {lvl}"]
            st.font.size = Pt(size)
            st.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
            st.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        for sec in self.doc.sections:
            sec.top_margin = Cm(2.5)
            sec.bottom_margin = Cm(2.5)
            sec.left_margin = Cm(2.8)
            sec.right_margin = Cm(2.8)

    # ---------- 基础片段 ----------
    def _feed_inline(self, p, text):
        r = InlineRenderer(p)
        r.feed(text)

    def _insert_image(self, raw_id, mime):
        path, _ = self.client.download_media(raw_id, mime) if self.client else (None, None)
        if not path or not os.path.exists(path):
            return
        try:
            width = Cm(MAX_IMG_WIDTH_CM)
            if HAS_PIL:
                try:
                    with Image.open(path) as im:
                        natural_cm = im.size[0] / 96.0 * 2.54
                    width = Cm(min(natural_cm, MAX_IMG_WIDTH_CM))
                except Exception:
                    pass
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=width)
        except Exception as e:
            print(f"  [警告] 图片插入失败 {raw_id}: {e}")

    def _insert_media_link(self, kind, raw_id, mime):
        path, _ = self.client.download_media(raw_id, mime) if self.client else (None, None)
        p = self.doc.add_paragraph()
        icon = "🔊 音频" if kind == "sound" else "🎬 视频"
        if path:
            rel = os.path.relpath(path, SCRIPT_DIR)
            self._add_hyperlink(p, rel, f"[{icon}] 点击打开（{os.path.basename(path)}）")
        else:
            p.add_run(f"[{icon}]（附件未能下载：{raw_id}）")

    @staticmethod
    def _add_hyperlink(paragraph, url, text):
        part = paragraph.part
        r_id = part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); rpr.append(color)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        run.append(rpr)
        t = OxmlElement("w:t"); t.text = text; run.append(t)
        link.append(run)
        paragraph._p.append(link)

    def _new_block_paragraph(self, kind, indent, align=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.75 * indent)
        pf.space_after = Pt(4)
        if align:
            p.alignment = align
        if kind == "quote":
            pf.left_indent = Cm(0.8)
            pf.right_indent = Cm(0.8)
        return p

    def _feed_tokens(self, p, tokens, kind, indent, align):
        """把一行的文本/媒体 token 写入段落；图片独占一段"""
        for typ, val in tokens:
            if typ == "text":
                self._feed_inline(p, val)
            elif typ == "img":
                p = self._new_block_paragraph(kind, indent, align)
                self._insert_image(val, self._cur_mime(val))
                p = self._new_block_paragraph(kind, indent, align)
            elif typ in ("sound", "video"):
                self._insert_media_link(typ, val, self._cur_mime(val))
        return p

    def _cur_mime(self, raw_id):
        return getattr(self, "_file_map", {}).get(raw_id, "")

    # ---------- 块级解析 ----------
    def render_body(self, content, file_map):
        self._file_map = file_map
        self.order_counter = 0
        for raw_line in content.split("\n"):
            line = raw_line.strip()
            if not line:
                continue
            self._render_line(line)

    def _render_line(self, line):
        # 分割线
        if re.match(r"^<hr\b[^>]*?/?>$", line):
            self._horizontal_rule()
            return
        # 裸标题标签（未包在 <text> 内的兜底情况）
        for tag, lvl in (("size", 3), ("mid-size", 4), ("h3-size", 5)):
            m = re.match(rf"^<{tag}>([\s\S]*?)</{tag}>$", line)
            if m:
                self._emit_heading(lvl, m.group(1))
                return
        # 复选框
        m = re.match(r"^<input\b([^>]*?)type=[\"']checkbox[\"']([^>]*?)\s*/?>(.*)$", line)
        if m:
            attrs = m.group(1) + m.group(2)
            indent = self._indent(attrs)
            checked = re.search(r"checked=[\"'](true|1)[\"']", attrs) is not None
            p = self._new_block_paragraph("checkbox", indent)
            mark = "☑ " if checked else "☐ "
            run = p.add_run(mark)
            run.bold = True
            self._feed_tokens(p, split_media(m.group(3)), "checkbox", indent, None)
            return
        # 有序列表
        m = re.match(r"^<order([^>]*)>([\s\S]*?)</order>$", line)
        if m:
            self.order_counter += 1
            indent = self._indent(m.group(1))
            p = self._new_block_paragraph("order", indent)
            p.add_run(f"{self.order_counter}. ")
            self._feed_tokens(p, split_media(m.group(2)), "order", indent, None)
            return
        # 无序列表
        m = re.match(r"^<bullet([^>]*)>([\s\S]*?)</bullet>$", line)
        if m:
            indent = self._indent(m.group(1))
            p = self._new_block_paragraph("bullet", indent)
            p.add_run("• ")
            self._feed_tokens(p, split_media(m.group(2)), "bullet", indent, None)
            return
        # 引用
        m = re.match(r"^<quote>([\s\S]*?)</quote>$", line)
        if m:
            p = self._new_block_paragraph("quote", 0)
            for run in self._style_quote():
                pass
            r0 = p.add_run("│ ")
            r0.font.color.rgb = RGBColor(0x9E, 0x9E, 0x9E)
            self._feed_tokens(p, split_media(m.group(1)), "quote", 0, None)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            return
        # 对齐
        m = re.match(r"^<align\b[^>]*align=[\"'](center|left|right)[\"'][^>]*>([\s\S]*?)</align>$", line)
        if m:
            amap = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT}
            self._emit_text(m.group(2), self._indent(""), amap[m.group(1)])
            return
        # 普通文本行
        m = re.match(r"^<text([^>]*)>([\s\S]*?)</text>$", line)
        if m:
            inner = m.group(2)
            if not inner.strip():
                return
            self._emit_text(inner, self._indent(m.group(1)), None)
            return
        # 兜底：去掉未知标签
        cleaned = re.sub(r"<[^>]+>", "", line)
        if cleaned.strip():
            p = self.doc.add_paragraph()
            self._feed_inline(p, cleaned)

    def _emit_heading(self, lvl, inner, align=None):
        p = self.doc.add_paragraph(style=f"Heading {lvl}")
        if align:
            p.alignment = align
        self._feed_tokens(p, split_media(inner), "heading", 0, align)

    def _emit_text(self, inner, indent, align):
        # 独占一行的标题标签
        for tag, lvl in (("size", 3), ("mid-size", 4), ("h3-size", 5)):
            m = re.match(rf"^<{tag}>([\s\S]*?)</{tag}>$", inner.strip())
            if m:
                self._emit_heading(lvl, m.group(1), align)
                return
        p = self._new_block_paragraph("text", indent, align)
        self._feed_tokens(p, split_media(inner), "text", indent, align)

    @staticmethod
    def _indent(attrs):
        m = re.search(r"indent=[\"'](\d+)[\"']", attrs or "")
        n = int(m.group(1)) if m else 0
        return max(0, n - 1)

    def _style_quote(self):
        return []

    def _horizontal_rule(self):
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "BFBFBF")
        pbdr.append(bottom)
        pPr.append(pbdr)

    # ---------- 结构 ----------
    def cover_and_toc(self, total, folder_count):
        t = self.doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.paragraph_format.space_before = Pt(120)
        run = t.add_run("小米笔记导出")
        run.font.size = Pt(32)
        run.bold = True
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
        sub = self.doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.add_run(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
                    f"共 {total} 篇笔记，{folder_count} 个分组").font.size = Pt(12)
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        # 目录域
        h = self.doc.add_paragraph(); h.add_run("目录").bold = True
        p = self.doc.add_paragraph()
        r = p.add_run()
        fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = 'TOC \\o "1-3" \\h \\z \\u'
        fs = OxmlElement("w:fldChar"); fs.set(qn("w:fldCharType"), "separate")
        tt = OxmlElement("w:t"); tt.text = "（在 Word 中按 Ctrl+A 后按 F9，或右键“更新域”即可生成目录）"
        fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end")
        r._r.append(fb); r._r.append(it); r._r.append(fs); r._r.append(tt); r._r.append(fe)
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def add_folder(self, name):
        h = self.doc.add_heading(name, level=1)
        return h

    def add_note(self, note):
        self.doc.add_heading(note["title"], level=2)
        meta = self.doc.add_paragraph()
        mr = meta.add_run(f"创建：{ts_to_str(note['create'])}    "
                          f"修改：{ts_to_str(note['modify'])}    分组：{note['folder_name']}")
        mr.font.size = Pt(9)
        mr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        self.render_body(note["content"], note["files"])
        self.doc.add_paragraph()  # 笔记间距

    def save(self, path):
        self.doc.save(path)


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #
def group_notes(notes):
    groups = {}
    for n in notes:
        groups.setdefault(n["folder_name"], []).append(n)
    key = "modify" if SORT_BY_MODIFY else "create"
    for lst in groups.values():
        lst.sort(key=lambda x: x.get(key) or 0)
    # “未分组”放最后
    names = sorted([g for g in groups if g != "未分组"])
    if "未分组" in groups:
        names.append("未分组")
    return names, groups


def build_and_save(notes, provider, filename_prefix="小米笔记导出"):
    print("④ 生成 Word（含图片处理）...")
    builder = DocBuilder(provider)
    names, groups = group_notes(notes)
    builder.cover_and_toc(len(notes), len(names))
    for name in names:
        builder.add_folder(name)
        for note in groups[name]:
            print(f"\r  渲染：{note['title'][:30]}", end="", flush=True)
            builder.add_note(note)
    print()
    out = os.path.join(SCRIPT_DIR,
                       f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
    builder.save(out)
    print(f"⑤ 完成！Word 已保存：\n   {out}")
    if os.path.isdir(ASSETS_DIR) and os.listdir(ASSETS_DIR):
        print(f"   附件目录：{ASSETS_DIR}")
    return out


class LocalJsonMediaProvider:
    """从浏览器导出的 JSON（base64 媒体）读取附件，无需联网、无需 Cookie"""
    def __init__(self, media):
        self.media = media or {}
        self._cache = {}

    def download_media(self, raw_id, mime):
        if raw_id in self._cache:
            return self._cache[raw_id]
        item = self.media.get(raw_id)
        if not item:
            self._cache[raw_id] = (None, None)
            return None, None
        m = re.match(r"data:([^;]+);base64,(.*)", item.get("data", ""), re.S)
        if not m:
            self._cache[raw_id] = (None, None)
            return None, None
        mt = (m.group(1) or item.get("mime") or mime or "").strip()
        cat = mt.split("/")[0].lower()
        kind = cat if cat in ("image", "audio", "video") else "image"
        ext = ext_from_mime(mt, "jpg" if kind == "image" else "bin")
        fid = raw_id.split(".", 1)[-1] if "." in raw_id else raw_id
        os.makedirs(ASSETS_DIR, exist_ok=True)
        path = os.path.join(ASSETS_DIR, f"{kind}_{fid[-12:]}.{ext}")
        if not os.path.exists(path):
            with open(path, "wb") as fp:
                fp.write(base64.b64decode(m.group(2)))
        result_kind = "img" if kind == "image" else kind
        self._cache[raw_id] = (path, result_kind)
        return path, result_kind


def find_json_file():
    p = os.path.join(SCRIPT_DIR, "mi_notes_data.json")
    if os.path.exists(p):
        return p
    dl = os.path.join(os.path.expanduser("~"), "Downloads", "mi_notes_data.json")
    if os.path.exists(dl):
        return dl
    return None


def run_from_json(json_path):
    if not json_path:
        json_path = find_json_file()
    if not json_path or not os.path.exists(json_path):
        print("未找到 mi_notes_data.json。请先在浏览器运行《①在浏览器控制台运行.txt》中的脚本，"
              "再把下载的 mi_notes_data.json 放到脚本文件夹（和 export_mi_notes.py 同目录）。")
        sys.exit(1)
    print(f"① 读取数据：{json_path}")
    with open(json_path, "r", encoding="utf-8") as fp:
        data = json.load(fp)
    folders = {str(k): (v or "未命名文件夹") for k, v in (data.get("folders") or {}).items()}
    media = data.get("media") or {}
    provider = LocalJsonMediaProvider(media)
    notes = []
    for n in data.get("notes", []):
        title = (n.get("title") or "").strip() or ts_to_str(n.get("createDate") or now_ms())
        folder_id = str(n.get("folderId") or "")
        files = {f.get("rawId"): f.get("mimeType", "") for f in n.get("files", []) if f.get("rawId")}
        notes.append({
            "id": str(n.get("id")),
            "title": title,
            "folder_id": folder_id,
            "folder_name": folders.get(folder_id) or "未分组",
            "content": n.get("content", ""),
            "files": files,
            "create": n.get("createDate") or 0,
            "modify": n.get("modifyDate") or 0,
        })
    print(f"② 共 {len(notes)} 篇笔记，{len(folders)} 个分组，媒体 {len(media)} 个。")
    if not notes:
        print("JSON 中没有任何笔记，请确认浏览器里能看到笔记、且脚本运行成功。")
        sys.exit(1)
    build_and_save(notes, provider)


def run_export(cookie):
    if requests is None:
        print("网络模式缺少 requests，请先运行：pip install -r requirements.txt")
        sys.exit(1)
    client = MiNoteClient(cookie)
    print("① 校验 Cookie ...")
    if not client.validate():
        print("Cookie 无效或已过期。建议改用浏览器抓取方式（见《①在浏览器控制台运行.txt》），"
              "或重新登录 i.mi.com 并复制最新 Cookie。")
        sys.exit(2)
    print("   Cookie 有效。")
    print("② 拉取笔记列表 ...")
    entries, folders = client.fetch_all_entries()
    print(f"   列表共 {len(entries)} 条，文件夹 {len(folders)} 个。")

    notes = []
    print("③ 逐篇拉取全文 ...")
    for i, entry in enumerate(entries, 1):
        nid = entry.get("id")
        try:
            detail = client.fetch_detail(nid)
            time.sleep(REQUEST_DELAY)
        except Exception as e:
            print(f"\n  [跳过] 详情获取失败 {nid}: {e}")
            detail = None
        note = normalize_note(entry, detail, folders)
        if note:
            notes.append(note)
        if i % 20 == 0:
            print(f"\r  已处理 {i}/{len(entries)}", end="", flush=True)
    print(f"\n   有效普通图文笔记 {len(notes)} 篇。")

    build_and_save(notes, client)


# --------------------------------------------------------------------------- #
# 演示模式（无需 Cookie，用于验证环境与排版）
# --------------------------------------------------------------------------- #
DEMO_NOTE = (
    "<text indent=\"1\">这是一篇<b>演示笔记</b>，用来验证 Word 排版效果。</text>\n"
    "<text indent=\"1\"></text>\n"
    "<text indent=\"1\"><size>大标题示例</size></text>\n"
    "<text indent=\"1\">支持<i>斜体</i>、<u>下划线</u>、<delete>删除线</delete>、"
    "<background color=\"#00FFFF\">黄色高亮（小米是BGR色值）</background>。</text>\n"
    "<text indent=\"1\"><mid-size>中标题与列表</mid-size></text>\n"
    "<bullet indent=\"1\">无序列表第一项</bullet>\n"
    "<bullet indent=\"2\">缩进的第二项</bullet>\n"
    "<order indent=\"1\">有序列表第一项</order>\n"
    "<order indent=\"1\">有序列表第二项</order>\n"
    "<input type=\"checkbox\" indent=\"1\" checked=\"true\" />已完成事项</input>\n"
    "<input type=\"checkbox\" indent=\"1\" checked=\"false\" />未完成事项</input>\n"
    "<quote>这是一段引用文字，用于演示引用块样式。</quote>\n"
    "<hr />\n"
    "<align align=\"center\">居中文字</align>\n"
    "<text indent=\"1\">图片、音频、视频会在真实导出时自动下载并插入/生成链接。</text>"
)


def run_demo():
    print("演示模式：不联网，使用内置样例生成 Word ...")
    builder = DocBuilder(None)
    builder.cover_and_toc(1, 1)
    builder.add_folder("演示分组")
    builder.add_note({
        "id": "demo", "title": "演示笔记", "folder_id": "", "folder_name": "演示分组",
        "content": DEMO_NOTE, "files": {},
        "create": now_ms(), "modify": now_ms(),
    })
    out = os.path.join(SCRIPT_DIR, "演示效果.docx")
    builder.save(out)
    print(f"演示 Word 已生成：{out}")


def main():
    ap = argparse.ArgumentParser(description="小米/红米笔记批量导出为 Word")
    ap.add_argument("--demo", action="store_true", help="使用内置样例生成演示 Word，无需登录")
    ap.add_argument("--from-json", dest="from_json", metavar="JSON路径", nargs="?", const="",
                    help="用浏览器脚本导出的 mi_notes_data.json 生成 Word（推荐，无需 Cookie）；"
                         "不填路径则自动在脚本文件夹/下载目录查找")
    args = ap.parse_args()

    if args.demo:
        run_demo()
        return

    if hasattr(args, "from_json") and args.from_json is not None:
        run_from_json(args.from_json or None)
        return

    if requests is None:
        print("缺少依赖 requests，请先运行：pip install -r requirements.txt")
        sys.exit(1)
    if not os.path.exists(COOKIE_FILE):
        print(f"未找到 cookie.txt。推荐改用浏览器抓取方式（见《①在浏览器控制台运行.txt》），"
              f"或在以下位置创建并粘贴 Cookie：\n  {COOKIE_FILE}")
        sys.exit(1)
    with open(COOKIE_FILE, "r", encoding="utf-8") as fp:
        cookie = fp.read().strip()
    # 兼容用户把 “Cookie: xxx” 整行粘贴进来
    if cookie.lower().startswith("cookie:"):
        cookie = cookie.split(":", 1)[1].strip()
    if "serviceToken" not in cookie:
        print("cookie.txt 里似乎缺少 serviceToken（登录态核心字段），请重新从浏览器复制完整 Cookie。")
        sys.exit(1)
    run_export(cookie)


if __name__ == "__main__":
    main()
